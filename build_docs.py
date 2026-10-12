# -*- coding: utf-8 -*-
"""把四个智能座舱 txt 原文分别转成精排 docx（不改字、不删内容）"""
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x3A, 0x6E)
BLUE = RGBColor(0x2E, 0x5E, 0xAA)
GOLD = RGBColor(0xC8, 0x9B, 0x3C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x55, 0x55, 0x55)
ALT = "F2F6FC"
MONO = "Consolas"
CN, CN_H = "宋体", "微软雅黑"

def setfont(run, cn=CN, size=10.5, bold=False, color=None, en="Times New Roman"):
    run.font.name = en
    run.font.size = Pt(size); run.font.bold = bold
    if color is not None: run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rPr.append(rf)
    rf.set(qn('w:eastAsia'), cn); rf.set(qn('w:ascii'), en); rf.set(qn('w:hAnsi'), en)

def shade(cell, hexc):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement('w:shd')
    s.set(qn('w:val'),'clear'); s.set(qn('w:color'),'auto'); s.set(qn('w:fill'),hexc)
    tcPr.append(s)

def cell_text(cell, txt, size=9.5, bold=False, color=None, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(str(txt)); setfont(r, size=size, bold=bold, color=color, cn=CN_H if bold else CN)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def new_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.3)
    s.left_margin = s.right_margin = Cm(2.4)
    return d

def cover(d, title, subtitle, meta_lines):
    for _ in range(4):
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("智能座舱多模态交互终端"); setfont(r, cn=CN_H, size=18, color=GREY)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); setfont(r, cn=CN_H, size=28, bold=True, color=NAVY)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle); setfont(r, cn=CN_H, size=14, color=BLUE)
    for _ in range(3): d.add_paragraph()
    t = d.add_table(rows=len(meta_lines), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,(k,v) in enumerate(meta_lines):
        c0,c1 = t.rows[i].cells
        shade(c0, "1F3A6E")
        cell_text(c0, k, 11, True, WHITE, True)
        cell_text(c1, v, 11)
        c0.width = Cm(4); c1.width = Cm(11)
    d.add_page_break()

def h1(d, text):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.page_break_before = True
    r = p.add_run(text); setfont(r, cn=CN_H, size=17, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr'); bd = OxmlElement('w:bottom')
    bd.set(qn('w:val'),'single'); bd.set(qn('w:sz'),'10'); bd.set(qn('w:space'),'4'); bd.set(qn('w:color'),'1F3A6E')
    b.append(bd); pPr.append(b)

def h2(d, text):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    r = p.add_run("▍"); setfont(r, cn=CN_H, size=13, bold=True, color=GOLD)
    r = p.add_run(text); setfont(r, cn=CN_H, size=13, bold=True, color=BLUE)

def h3(d, text):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); setfont(r, cn=CN_H, size=11, bold=True, color=NAVY)

def body(d, text, indent=True, size=10.5):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(3); pf.line_spacing = 1.4
    if indent: pf.first_line_indent = Pt(size*2)
    r = p.add_run(text); setfont(r, size=size)

def bullet(d, text, size=10.5):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Pt(20); pf.space_after = Pt(2); pf.line_spacing = 1.35
    r = p.add_run("• "); setfont(r, size=size, bold=True, color=BLUE)
    r = p.add_run(text); setfont(r, size=size)

def mono_block(d, lines):
    """代码/目录树块：浅灰底、等宽字体"""
    for ln in lines:
        p = d.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Pt(14); pf.space_after = Pt(0); pf.line_spacing = 1.15
        r = p.add_run(ln if ln else " ")
        setfont(r, size=9, en=MONO, cn=CN)
        # 浅灰底
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),'F4F5F7')
        pPr.append(shd)

def md_table(d, rows):
    """rows: [[c1,c2,...], ...]，首行表头"""
    ncol = max(len(r) for r in rows)
    rows = [r + [""]*(ncol-len(r)) for r in rows]
    t = d.add_table(rows=len(rows), cols=ncol)
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j,htxt in enumerate(rows[0]):
        c = t.rows[0].cells[j]; shade(c,"1F3A6E")
        cell_text(c, htxt.strip(), 9.5, True, WHITE, True)
    for i,rw in enumerate(rows[1:]):
        for j,v in enumerate(rw):
            c = t.rows[i+1].cells[j]
            if i % 2 == 1: shade(c, ALT)
            cell_text(c, v.strip(), 9.5, center=(j==0))
    d.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------- 解析 ----------
def is_sep_line(s, ch='='):
    t = s.strip()
    return len(t) >= 10 and set(t) == {ch}

def is_table_row(s):
    t = s.strip()
    return t.startswith('|') and t.endswith('|') and t.count('|') >= 2

def split_md_row(s):
    return [c.strip() for c in s.strip().strip('|').split('|')]

def looks_mono(ln):
    t = ln.strip()
    return (ln.startswith('  ') and t and not t.endswith('。') and not t.endswith('：')
            and any(k in ln for k in ['├','└','│','/','.py','.js','.json','.env','$','cd ','npm ','python ',
                                       'GET ','POST ','PUT ','WS ',':8000',':7861',':5173',':8787',
                                       'http','端口','端口','VITE_','DEEPSEEK','AMAP'])
            and len(t) < 80)

def convert(txt_path, out_docx, cover_title, cover_sub, meta):
    lines = open(txt_path, encoding='utf-8').read().splitlines()
    d = new_doc()
    cover(d, cover_title, cover_sub, meta)

    i = 0
    n = len(lines)
    mono_buf = []
    def flush_mono():
        nonlocal mono_buf
        if mono_buf:
            mono_block(d, mono_buf); mono_buf = []

    while i < n:
        raw = lines[i]
        s = raw.rstrip()
        st = s.strip()

        # 分隔线 =（成对出现：章节标题上下各一条）
        if is_sep_line(s, '='):
            j = i+1
            while j < n and not lines[j].strip(): j += 1
            if j < n:
                nxt = lines[j].strip()
                # 下一个非空行是 ## / ### markdown 标题 → 这是章节下边框，跳过
                if nxt.startswith('## ') or nxt.startswith('### '):
                    i = j; continue
                # 否则是章节大标题
                h1(d, nxt)
                i = j+1
                continue
            i += 1; continue

        # 小分隔线 ---
        if is_sep_line(s, '-') and len(st) >= 8:
            # 上一行非空则上一行是小节标题
            # 向前找上一非空行（在缓冲里），这里简化：上一条已入文档，跳过
            i += 1; continue

        # ## 标题
        if st.startswith('## '):
            flush_mono(); h2(d, st[3:].strip()); i += 1; continue
        if st.startswith('### '):
            flush_mono(); h3(d, st[4:].strip()); i += 1; continue

        # Markdown 表格
        if is_table_row(s):
            flush_mono()
            tbl = []
            while i < n and is_table_row(lines[i].strip()):
                row = split_md_row(lines[i])
                # 跳过分隔行 |---|---|
                if not all(set(c) <= set('-: ') for c in row):
                    tbl.append(row)
                i += 1
            if tbl: md_table(d, tbl)
            continue

        # 空行
        if not st:
            flush_mono(); i += 1; continue

        # 项目符号
        if st.startswith('· ') or st.startswith('- '):
            flush_mono(); bullet(d, st[2:].strip()); i += 1; continue

        # 编号项  1. / (1)
        if re.match(r'^\d+[\.\、]', st) or st.startswith('('):
            flush_mono(); bullet(d, st); i += 1; continue

        # 等宽块（代码/目录树/端口表）
        if looks_mono(raw):
            mono_buf.append(st)
            i += 1; continue
        else:
            flush_mono()

        # 普通段落
        body(d, st)
        i += 1

    flush_mono()
    d.save(out_docx)
    print("SAVED", out_docx)

BASE = r"C:\Users\ASUS\Desktop\smart_seat"
jobs = [
    ("智能座舱项目综合总结.txt",
     "智能座舱项目01-项目综合总结.docx",
     "项 目 综 合 总 结",
     "Project Comprehensive Summary",
     [("项目全称","基于 LoongArch 端侧 AI 的智能座舱多模态交互终端"),
      ("文档版本","V1.0    2026 年 9 月")]),
    ("智能座舱项目-用户手册.txt",
     "智能座舱项目02-用户手册.docx",
     "用 户 手 册",
     "Smart Cabin Assistant User Manual",
     [("适用对象","终端用户 / 演示人员 / 评测人员"),
      ("文档版本","V1.0    2026-09-10")]),
    ("智能座舱项目-功能测试报告.txt",
     "智能座舱项目03-功能测试报告.docx",
     "功 能 测 试 报 告",
     "Functional Test Report",
     [("文档编号","SC-TEST-2026-001"),
      ("测试类型","功能 / 接口 / WebSocket / 兼容性"),
      ("测试结论","通过（详见正文）")]),
    ("智能座舱项目-其他相关文档.txt",
     "智能座舱项目04-其他相关文档.docx",
     "其 他 相 关 文 档",
     "Deployment · API · Data Dictionary · O&M",
     [("涵盖范围","部署运维 / API / 数据字典 / 硬件 / 术语 / 安全 / 版本 / 排障"),
      ("文档版本","V1.0    2026-09-10")]),
]

for src, dst, t1, t2, meta in jobs:
    convert(os.path.join(BASE, src), os.path.join(BASE, dst), t1, t2, meta)
